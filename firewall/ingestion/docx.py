"""Word documents: body/tables/headers/footers, hidden runs (vanish, white, tiny font),
reviewer comments and core properties."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import RGBColor

from ..core.context import ScanContext
from ..core.models import Channel, Segment, Source
from .base import as_bytes

WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _run_hidden_reason(run) -> str | None:
    font = run.font
    if font.hidden:
        return "hidden (vanish)"
    try:
        if font.color is not None and font.color.rgb == WHITE:
            return "white text"
    except (AttributeError, ValueError):
        pass
    if font.size is not None and font.size.pt < 2:
        return f"{font.size.pt:g}pt font"
    return None


def extract_docx(data: bytes | str, ctx: ScanContext) -> list[Segment]:
    doc = Document(io.BytesIO(as_bytes(data)))
    lines: list[str] = []
    extra: list[Segment] = []

    def paragraph(p, loc: str) -> None:
        visible = []
        for run in p.runs:
            if not run.text:
                continue
            if reason := _run_hidden_reason(run):
                extra.append(Segment(text=run.text, channel=Channel.HIDDEN, location=f"{loc} ({reason})", emit=False))
            else:
                visible.append(run.text)
        lines.append("".join(visible) if p.runs else p.text)

    for p in doc.paragraphs:
        paragraph(p, "docx:body")
    for t, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    paragraph(p, f"docx:table{t}")
    for s, section in enumerate(doc.sections):
        for part, name in ((section.header, "header"), (section.footer, "footer")):
            for p in part.paragraphs:
                if p.text.strip():
                    paragraph(p, f"docx:{name}{s}")

    props = doc.core_properties
    for field in ("title", "subject", "keywords", "comments", "category", "author", "last_modified_by"):
        val = getattr(props, field, None)
        if val and len(str(val).strip()) >= 3:
            extra.append(Segment(text=str(val), channel=Channel.METADATA, location=f"docx:property:{field}", emit=False))

    for part in doc.part.package.iter_parts():
        if str(part.partname).endswith("/comments.xml"):
            xml = part.blob.decode("utf-8", "replace")
            for c in re.findall(r"<w:comment\b.*?</w:comment>", xml, re.S):
                text = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", c))
                if text.strip():
                    extra.append(Segment(text=text, channel=Channel.COMMENT, location="docx:comment", emit=False))

    body = "\n".join(line for line in lines if line.strip())
    return [Segment(text=body, channel=Channel.VISIBLE, location="docx:body")] + extra
